from __future__ import annotations

import io
from pathlib import Path

from docx import Document
from pypdf import PdfReader

from server.app.models.schemas import ResumeProfile


class ResumeValidationError(ValueError):
    """A user-correctable resume upload problem."""


def extract_resume_text(filename: str, content: bytes, max_bytes: int) -> str:
    if not filename:
        raise ResumeValidationError("Please choose a PDF or DOCX resume.")
    extension = Path(filename).suffix.lower()
    if extension not in {".pdf", ".docx"}:
        raise ResumeValidationError("That file type isn't supported. Please upload a PDF or DOCX resume.")
    if len(content) > max_bytes:
        raise ResumeValidationError("That resume is over 5 MB. Please upload a smaller file.")
    if not content:
        raise ResumeValidationError("The uploaded file is empty. Please choose another resume.")

    try:
        if extension == ".pdf":
            reader = PdfReader(io.BytesIO(content))
            text = "\n".join(page.extract_text() or "" for page in reader.pages)
        else:
            document = Document(io.BytesIO(content))
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_text = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            text = "\n".join((*paragraphs, *table_text))
    except Exception as exc:
        raise ResumeValidationError(
            "I couldn't read that resume. Please upload a valid, text-based PDF or DOCX file."
        ) from exc

    text = "\n".join(line.strip() for line in text.splitlines() if line.strip())
    if len(text) < 100:
        raise ResumeValidationError(
            "I couldn't find enough readable text in that resume. Scanned PDFs aren't supported yet; "
            "please upload a text-based PDF or DOCX."
        )
    return text


def validate_resume_profile(data: dict[str, object]) -> ResumeProfile:
    return ResumeProfile.model_validate(data)
